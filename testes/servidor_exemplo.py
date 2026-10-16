r"""
Servidor local de EXEMPLO que faz o papel da Central de Gestão (API das
vistorias e fotos) e da Intranet (requisição do Banco D com a planilha do
banco anexada), com as propostas fictícias de dados_exemplo.py.

Existe porque esta é a versão de portfólio: os endereços reais foram
trocados por *.example.com. Com ele, o programa roda de ponta a ponta sem
nenhum sistema da empresa - nos testes e numa demonstração:

    python testes\servidor_exemplo.py            (mostra as variáveis pra apontar o programa pra ele)

Os testes (test_unidade.py e regressao.py) sobem o servidor sozinhos numa
porta livre. Só a biblioteca padrão + Pillow (fotos) + openpyxl (planilha).

Rotas (as mesmas que o programa usa - ver central_gestao.py e intranet.py):
  Central de Gestão:  POST /auth, GET /auth/validToken, GET /inspection?id=,
                      GET /inspection/<_id>, GET /public/images/...
  Intranet:           GET /login, POST /login/username, POST /login/senha,
                      GET /portal/requisicao[/busca?descricao=],
                      GET /portal/requisicao/entrada/detalhar/<n>,
                      GET /portal/requisicao/download/id/<n>
"""

import hashlib
import io
import json
import os
import re
import sys
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import parse_qs, unquote, urlparse

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import dados_exemplo  # noqa: E402
from central_gestao import API_KEY_LOGIN  # noqa: E402

COOKIE = "SESSAO_INTRANET"
_fotos_feitas = {}
_trava = threading.Lock()


# ---------------------------------------------------------------- fotos e planilha fictícias

def foto(caminho):
    """JPEG 1280x960 desenhado na hora: um "cômodo" com o nome do ambiente e
    o carimbo de data/GPS no canto, como as fotos do app de vistoria. Mesmo
    caminho = mesma foto (o teste de regressão compara resultados)."""
    from PIL import Image, ImageDraw, ImageFont

    with _trava:
        if caminho in _fotos_feitas:
            return _fotos_feitas[caminho]
    semente = hashlib.md5(caminho.encode()).digest()
    cor = (90 + semente[0] % 120, 90 + semente[1] % 120, 90 + semente[2] % 120)
    nome = os.path.basename(caminho).split("_")[0]
    numero = re.search(r"\[(\d+)\]", caminho)
    im = Image.new("RGB", (1280, 960), cor)
    d = ImageDraw.Draw(im)
    d.polygon([(0, 960), (1280, 960), (980, 640), (300, 640)], fill=tuple(c - 40 for c in cor))  # piso
    d.rectangle((420 + semente[3] % 200, 180, 760 + semente[3] % 200, 460), fill=(225, 238, 250),
                outline=(255, 255, 255), width=10)                                             # janela
    try:
        fonte, pequena = ImageFont.truetype("arial.ttf", 84), ImageFont.truetype("arial.ttf", 30)
    except OSError:
        fonte, pequena = ImageFont.load_default(size=84), ImageFont.load_default(size=30)
    d.text((640, 560), f"{nome} {int(numero.group(1)) + 1 if numero else ''}".strip(), font=fonte,
           fill="white", anchor="mm", stroke_width=3, stroke_fill="black")
    d.text((1250, 930), "14/10/2026 10:30  -23.000000, -46.000000  (foto de exemplo)", font=pequena,
           fill="yellow", anchor="rs", stroke_width=2, stroke_fill="black")
    saida = io.BytesIO()
    im.save(saida, "JPEG", quality=85)
    with _trava:
        _fotos_feitas[caminho] = saida.getvalue()
    return saida.getvalue()


def planilha_banco_d(proposta):
    """A planilha que o Banco D manda (só as células que o programa lê)."""
    from openpyxl import Workbook

    d = dados_exemplo.PLANILHAS_BANCO_D[proposta]
    wb = Workbook()
    ws = wb.active
    ws.title = "FORMULARIO"
    ws["E5"], ws["F5"] = "Junção:", d["juncao"]
    ws["E6"], ws["F6"] = "Nome da junção:", d["nome_juncao"]
    ws["E7"], ws["F7"] = "Protocolo:", proposta
    ws["E36"], ws["F36"] = "Tipo de operação:", d["tipo"]
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


# ---------------------------------------------------------------- HTTP

class _Tratador(BaseHTTPRequestHandler):
    propostas = None  # {nº: vistoria}, montado no iniciar()

    def log_message(self, *args):  # silencioso (os testes mostram o que importa)
        pass

    # ---- respostas
    def _enviar(self, status, corpo=b"", tipo="text/html; charset=utf-8", cabecalhos=None):
        if isinstance(corpo, str):
            corpo = corpo.encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", tipo)
        self.send_header("Content-Length", str(len(corpo)))
        for k, v in (cabecalhos or {}).items():
            self.send_header(k, v)
        self.end_headers()
        self.wfile.write(corpo)

    def _json(self, dados, status=200):
        self._enviar(status, json.dumps(dados, ensure_ascii=False), "application/json; charset=utf-8")

    def _redirecionar(self, para, cabecalhos=None):
        self._enviar(302, b"", cabecalhos={"Location": para, **(cabecalhos or {})})

    def _corpo(self):
        tamanho = int(self.headers.get("Content-Length") or 0)
        return self.rfile.read(tamanho) if tamanho else b""

    def _logado_api(self):
        return self.headers.get("authorization") == "Barear " + dados_exemplo.TOKEN

    def _logado_intranet(self):
        return f"{COOKIE}=ok" in (self.headers.get("Cookie") or "")

    # ---- rotas
    def do_GET(self):
        url = urlparse(self.path)
        consulta = parse_qs(url.query)
        caminho = unquote(url.path)
        # Central de Gestão
        if caminho == "/auth/validToken":
            return self._enviar(200 if self._logado_api() else 401,
                                "VALID TOKEN" if self._logado_api() else "INVALID", "text/plain")
        if caminho == "/inspection":
            if not self._logado_api():
                return self._json({"error": "unauthorized"}, 401)
            busca = (consulta.get("id") or [""])[0]
            if not busca.isdigit():  # como na API real: texto dá erro 500
                return self._json({"error": "CastError"}, 500)
            v = self.propostas.get(busca)
            achadas = [{"id": v["id"], "_id": v["_id"], "fotos": "id-das-fotos"}] if v else []
            return self._json({"result": achadas, "totalFound": len(achadas)})
        m = re.fullmatch(r"/inspection/(\w+)", caminho)
        if m:
            if not self._logado_api():
                return self._json({"error": "unauthorized"}, 401)
            v = next((v for v in self.propostas.values() if v["_id"] == m.group(1)), None)
            return self._json(v) if v else self._json({"error": "not found"}, 404)
        if caminho.startswith("/public/images/"):  # públicas (sem token), JPEG mesmo com .png
            existe = any(caminho.lstrip("/") in img for v in self.propostas.values()
                         for item in v["fotos"]["fotos"] for img in item["images"])
            return self._enviar(200, foto(caminho), "image/png") if existe else self._enviar(404, b"not found")
        # Intranet
        if caminho == "/login":
            return self._enviar(200, '<form action="/login/username" method="post">'
                                     '<input type="hidden" name="_csrf" value="exemplo">'
                                     '<input type="text" name="username"><button>Continuar</button></form>')
        if caminho.startswith("/portal/") and not self._logado_intranet():
            return self._redirecionar("/login")
        if caminho == "/portal/requisicao":
            return self._enviar(200, "<h1>Requisições</h1>")
        if caminho == "/portal/requisicao/busca":
            proposta = (consulta.get("descricao") or [""])[0]
            links = "".join(f'<li><a href="/portal/requisicao/entrada/detalhar/{p}">{p} - Laudo - BANCO D</a></li>'
                            for p in dados_exemplo.PLANILHAS_BANCO_D if p == proposta)
            return self._enviar(200, f"<ul>{links}</ul>")
        m = re.fullmatch(r"/portal/requisicao/entrada/detalhar/(\d+)", caminho)
        if m and m.group(1) in dados_exemplo.PLANILHAS_BANCO_D:
            return self._enviar(200, f'<h2>Arquivos ativos</h2><a href="/portal/requisicao/download/id/{m.group(1)}">'
                                     f'Formulario_Banco_D - 14-10-2026.xlsx</a>'
                                     f'<a href="/portal/requisicao/download/id/0">foto_extra.jpg</a>')
        m = re.fullmatch(r"/portal/requisicao/download/id/(\d+)", caminho)
        if m and m.group(1) in dados_exemplo.PLANILHAS_BANCO_D:
            return self._enviar(200, planilha_banco_d(m.group(1)),
                                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        return self._enviar(404, b"not found")

    def do_POST(self):
        caminho = urlparse(self.path).path
        corpo = self._corpo()
        if caminho == "/auth":
            try:
                dados = json.loads(corpo or b"{}")
            except ValueError:
                dados = {}
            if self.headers.get("x-api-key") != API_KEY_LOGIN:
                return self._json({"auth": False, "message": "api key"}, 403)
            if (dados.get("email"), dados.get("password")) == (dados_exemplo.USUARIO, dados_exemplo.SENHA):
                return self._json({"auth": True, "token": dados_exemplo.TOKEN})
            return self._json({"auth": False, "message": "Usuário ou senha inválidos"}, 401)
        form = {k: v[0] for k, v in parse_qs(corpo.decode("utf-8")).items()}
        if caminho == "/login/username":  # 1ª etapa: só o usuário; devolve o formulário da senha
            if form.get("username") != dados_exemplo.USUARIO:
                return self._redirecionar("/login")
            return self._enviar(200, f'<form action="/login/senha" method="post">'
                                     f'<input type="hidden" name="username" value="{form["username"]}">'
                                     f'<input type="password" name="password">'
                                     f'<input type="checkbox" name="lembrar" value="1"><button>Entrar</button></form>')
        if caminho == "/login/senha":
            if (form.get("username"), form.get("password")) == (dados_exemplo.USUARIO, dados_exemplo.SENHA):
                return self._redirecionar("/portal/requisicao", {"Set-Cookie": f"{COOKIE}=ok; Path=/"})
            return self._redirecionar("/login")
        return self._enviar(404, b"not found")


def iniciar(porta=0):
    """Sobe o servidor numa thread (porta 0 = uma livre). Devolve (servidor, url)."""
    _Tratador.propostas = dados_exemplo.propostas()
    servidor = ThreadingHTTPServer(("127.0.0.1", porta), _Tratador)
    threading.Thread(target=servidor.serve_forever, daemon=True).start()
    return servidor, f"http://127.0.0.1:{servidor.server_address[1]}"


def variaveis(url):
    """Variáveis de ambiente que apontam o programa pro servidor de exemplo."""
    return {"CENTRAL_GESTAO_API": url, "INTRANET_URL": url,
            "CENTRAL_GESTAO_USUARIO": dados_exemplo.USUARIO, "CENTRAL_GESTAO_SENHA": dados_exemplo.SENHA,
            "INTRANET_USUARIO": dados_exemplo.USUARIO, "INTRANET_SENHA": dados_exemplo.SENHA}


if __name__ == "__main__":
    servidor, url = iniciar(int(sys.argv[1]) if len(sys.argv) > 1 else 8765)
    print(f"Servidor de exemplo em {url} - propostas: {', '.join(dados_exemplo.propostas())}")
    print("Para o programa usar este servidor (PowerShell, antes de rodar python app.py):")
    for k, v in variaveis(url).items():
        print(f'  $env:{k} = "{v}"')
    print("Ctrl+C encerra.")
    try:
        threading.Event().wait()
    except KeyboardInterrupt:
        servidor.shutdown()
