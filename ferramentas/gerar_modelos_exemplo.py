r"""
Gera os modelos de Excel FICTÍCIOS desta versão de portfólio:

    modelos\banco_c\modelo_comparativo.xlsm, modelo_evolutivo.xlsm
    modelos\banco_d\modelo_<tipo de operação>.xlsx (5)

Os modelos de verdade são dos bancos e ficaram fora do repositório. Estes
reproduzem só o que o programa usa - nomes das abas, blocos de foto
mesclados nas linhas/colunas medidas em perfis.py, títulos dos quadros,
células do formulário que formulario.py preenche, área e escala de
impressão, quebras de página e as células escritas "Logo" - com textos
genéricos. Assim o programa roda de ponta a ponta e os testes passam.

    python ferramentas\gerar_modelos_exemplo.py

Só openpyxl (não precisa de Excel). Rodar de novo regrava os arquivos; se
mudar algo aqui, rode também testes\regressao.py --gravar (no Windows, com
Excel) pra atualizar os gabaritos.

Detalhe do .xlsm: o openpyxl grava um arquivo sem macros com o tipo de
conteúdo de .xlsx, e o Excel recusa abrir .xlsm assim ("formato ou extensão
inválido"). Depois de salvar, o tipo do workbook em [Content_Types].xml é
trocado pro de pasta habilitada para macro (sem vbaProject.bin, que é
opcional) - é o mesmo arquivo que o Excel gravaria com "Salvar como .xlsm".
"""

import os
import sys
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, REPO)
from perfis import BANCO_C, BANCO_D  # noqa: E402

FINO = Side(style="thin", color="7F7F7F")
CABECALHO = PatternFill("solid", fgColor="D9E1F2")   # azul-acinzentado claro (títulos)
ROTULO = PatternFill("solid", fgColor="EDEDED")      # cinza (rótulos dos quadros)
ENTRADA = PatternFill("solid", fgColor="FFF9E5")     # amarelo bem claro (campos a preencher)
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
# rótulos sem quebra de linha: célula com quebra faz o Excel aumentar a altura da linha sozinho
# e a página deixa de caber na folha; nas áreas mescladas de texto longo (comentário), com quebra
ESQUERDA = Alignment(horizontal="left", vertical="center")
TEXTO_LONGO = Alignment(horizontal="left", vertical="top", wrap_text=True)
XLSM_TIPO = "application/vnd.ms-excel.sheet.macroEnabled.main+xml"
XLSX_TIPO = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"


# ---------------------------------------------------------------- utilitários

def moldura(ws, faixa):
    """Borda fina em volta de `faixa` (célula ou intervalo)."""
    c1, l1, c2, l2 = range_boundaries(faixa)
    for linha in range(l1, l2 + 1):
        for col in range(c1, c2 + 1):
            ws.cell(linha, col).border = Border(
                left=FINO if col == c1 else None, right=FINO if col == c2 else None,
                top=FINO if linha == l1 else None, bottom=FINO if linha == l2 else None)


def mesclar(ws, faixa, valor=None, negrito=False, fundo=None, alinhar=CENTRO, borda=True, tamanho=None):
    """Mescla, escreve no canto e formata."""
    if ":" in faixa:
        ws.merge_cells(faixa)
    canto = ws[faixa.split(":")[0]]
    if valor is not None:
        canto.value = valor
    canto.font = Font(bold=negrito, size=tamanho or 10)
    canto.alignment = alinhar
    if fundo:
        c1, l1, c2, l2 = range_boundaries(faixa if ":" in faixa else f"{faixa}:{faixa}")
        for linha in range(l1, l2 + 1):
            for col in range(c1, c2 + 1):
                ws.cell(linha, col).fill = fundo
    if borda:
        moldura(ws, faixa if ":" in faixa else f"{faixa}:{faixa}")
    return canto


def rotulo(ws, celula, texto, negrito=False):
    c = ws[celula]
    c.value = texto
    c.font = Font(bold=negrito, size=10)
    c.alignment = ESQUERDA
    return c


def campo(ws, faixa, formato=None, valor=None):
    """Campo a preencher (fundo claro + borda)."""
    canto = mesclar(ws, faixa, valor, fundo=ENTRADA, alinhar=ESQUERDA)
    if formato:
        canto.number_format = formato
    return canto


def impressao(ws, area, escala=None, ajustar_largura=None, ajustar_altura=None, quebras=(), margens=None,
              paisagem=False, ajustar_pagina=False):
    """Área de impressão, escala, A4 e quebras manuais (id = última linha
    antes da quebra). `ajustar_largura/altura` só valem com `ajustar_pagina`
    ("ajustar a N páginas"); sem ele, vale a escala - como nos modelos reais,
    que guardam os dois."""
    ws.print_area = area
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if paisagem else "portrait"
    if escala:
        ws.page_setup.scale = escala
    if ajustar_largura is not None:
        ws.page_setup.fitToWidth = ajustar_largura
    if ajustar_altura is not None:
        ws.page_setup.fitToHeight = ajustar_altura
    if ajustar_pagina:
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    for linha in quebras:
        ws.row_breaks.append(Break(id=linha))
    esq, dir_, topo, base = margens or (0.51, 0.51, 0.79, 0.79)
    ws.page_margins.left, ws.page_margins.right = esq, dir_
    ws.page_margins.top, ws.page_margins.bottom = topo, base


def larguras(ws, colunas):
    for col, largura in colunas.items():
        ws.column_dimensions[col].width = largura


def alturas(ws, linhas):
    for linha, altura in linhas.items():
        ws.row_dimensions[linha].height = altura


def logo(ws, faixa="A1:B2"):
    """Célula escrita só "Logo": o programa troca pela logo da empresa."""
    mesclar(ws, faixa, "Logo", borda=False)


def salvar(wb, caminho):
    os.makedirs(os.path.dirname(caminho), exist_ok=True)
    wb.save(caminho)
    if caminho.lower().endswith(".xlsm"):
        _marcar_como_xlsm(caminho)
    print("gerado:", os.path.relpath(caminho, REPO))


def _marcar_como_xlsm(caminho):
    """Troca o tipo do workbook em [Content_Types].xml pro de .xlsm (ver o topo)."""
    with zipfile.ZipFile(caminho) as z:
        itens = [(info, z.read(info.filename)) for info in z.infolist()]
    tmp = caminho + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for info, dados in itens:
            if info.filename == "[Content_Types].xml":
                dados = dados.replace(XLSX_TIPO.encode(), XLSM_TIPO.encode())
            z.writestr(info, dados)
    os.replace(tmp, caminho)


# ---------------------------------------------------------------- Banco C

def _aba_fotos_banco_c(wb):
    """'Fotos internas imóvel': páginas de 50 linhas (a 1ª com o título), 3
    fileiras de 2 fotos, bloco de 13 linhas e "Descrição:" embaixo - as
    medidas de perfis.BANCO_C.fotos."""
    lay = BANCO_C.fotos
    ws = wb.create_sheet(lay.nome)
    larguras(ws, {"A": 2.0, "B": 9.62, "C": 6.38, "G": 9.62, "H": 3.0, "I": 9.62, "J": 6.38, "N": 9.62, "O": 2.0})
    mesclar(ws, "B2:N2", "REGISTRO FOTOGRÁFICO - ÁREAS PRIVATIVAS DO IMÓVEL", negrito=True, fundo=CABECALHO,
            tamanho=12)
    alturas(ws, {1: 10.5, 2: 21.0})
    for pagina in range(lay.paginas_no_modelo):
        base = pagina * lay.linhas_por_pagina
        alturas(ws, {base + 3: 10.5, base + 51: 10.5, base + 52: 10.5})
        for n, linha in enumerate(lay.linhas_foto):
            ini = base + linha
            desc = ini + lay.altura_bloco
            alturas(ws, {desc: 15.75, desc + 1: 15.75, desc + 2: 10.5})
            for col_ini, col_fim, col_desc in lay.colunas:
                mesclar(ws, f"{col_ini}{ini}:{col_fim}{ini + lay.altura_bloco - 1}")
                rotulo_desc = ws[f"{col_ini}{desc}"]
                rotulo_desc.value = "Descrição:"
                rotulo_desc.font = Font(size=9)
                rotulo_desc.fill = ROTULO
                mesclar(ws, f"{col_desc}{desc}:{col_fim}{desc + 1}", alinhar=ESQUERDA)
    # o modelo vem com área de impressão e quebras só nas 4 primeiras páginas
    # (como os modelos reais): o programa estende conforme o nº de fotos
    impressao(ws, "A1:O202", escala=85, ajustar_altura=0, quebras=(52, 102, 152, 202),
              margens=(0.0, 0.0, 0.35, 0.35))
    return ws


def _aba_locais_banco_c(wb):
    """'Mapa e fotos locais': mapa + bloco "REGISTRO FOTOGRÁFICO LOCALIDADE"
    (linhas 29-61) com os 4 quadros de perfis.BANCO_C.locais."""
    loc = BANCO_C.locais
    ws = wb.create_sheet(loc.nome)
    larguras(ws, {"A": 1.88, "B": 6.25, "H": 3.0, "O": 1.88})
    alturas(ws, {1: 10.5, 2: 21.0, 3: 10.35, 24: 15.75, 25: 15.75, 26: 15.75, 27: 15.75, 28: 10.35, 30: 10.35,
                 45: 10.35, 46: 10.35, 61: 10.35})
    mesclar(ws, "B2:N2", "MAPA COM A LOCALIZAÇÃO DO IMÓVEL", negrito=True, fundo=CABECALHO, tamanho=12)
    mesclar(ws, "B4:N23", "(mapa da região - inserido pela equipe)")
    ws["B4"].font = Font(italic=True, color="7F7F7F")
    mesclar(ws, "B24:C27", "Comentários sobre a região:", fundo=ROTULO)
    mesclar(ws, "D24:N27", alinhar=ESQUERDA)
    ini, _ = loc.bloco_molde
    mesclar(ws, f"B{ini}:N{ini}", "REGISTRO FOTOGRÁFICO LOCALIDADE", negrito=True, fundo=CABECALHO)
    for q in loc.quadros:
        mesclar(ws, q.bloco)
        c1, _, c2, l2 = range_boundaries(q.bloco)
        mesclar(ws, f"{get_column_letter(c1)}{l2 + 1}:{get_column_letter(c2)}{l2 + 1}", q.titulo, negrito=True,
                fundo=ROTULO)
    impressao(ws, f"A1:{loc.col_fim_impressao}61", escala=83, ajustar_largura=0, margens=(0.0, 0.0, 0.35, 0.35))
    return ws


def _aba_resumo_banco_c(wb, tipo):
    ws = wb.active
    ws.title = "Resumo"
    larguras(ws, {"A": 3.0, "B": 22.0, "C": 18.0, "D": 18.0, "E": 18.0, "F": 18.0, "G": 3.0})
    alturas(ws, {1: 20.0, 2: 20.0, 3: 20.0})
    logo(ws, "B1:B3")
    mesclar(ws, "C1:F3", f"LAUDO DE AVALIAÇÃO - BANCO C (MODELO DE EXEMPLO)\nMétodo {tipo}", negrito=True,
            fundo=CABECALHO, tamanho=12)
    campos = ["Proposta", "Proponente", "Endereço", "Bairro / Cidade / UF", "Tipo do imóvel", "Área privativa (m²)",
              "Área total (m²)", "Data da vistoria", "Valor de mercado (R$)", "Valor de liquidação forçada (R$)"]
    for n, nome in enumerate(campos):
        linha = 5 + n * 2
        rotulo(ws, f"B{linha}", nome + ":", negrito=True)
        campo(ws, f"C{linha}:F{linha}")
    impressao(ws, "A1:G26", escala=90)


def _abas_extras_banco_c(wb, tipo):
    """Abas de cálculo, diferentes entre os dois modelos (só pra existir e
    imprimir; o programa não mexe nelas). Uma oculta: o programa não pode
    tentar pôr logo em aba oculta."""
    if tipo == "Comparativo":
        ws = wb.create_sheet("Amostras")
        larguras(ws, {"A": 3.0, "B": 28.0, "C": 14.0, "D": 14.0, "E": 14.0, "F": 14.0})
        mesclar(ws, "B2:F2", "AMOSTRAS DE MERCADO (EXEMPLO)", negrito=True, fundo=CABECALHO)
        for n, nome in enumerate(["Endereço", "Área (m²)", "Valor (R$)", "R$/m²", "Fonte"]):
            mesclar(ws, f"{get_column_letter(2 + n)}4", nome, negrito=True, fundo=ROTULO)
        for linha in range(5, 15):
            ws[f"E{linha}"] = f'=IFERROR(D{linha}/C{linha},"")'
            for col in "BCDEF":
                moldura(ws, f"{col}{linha}")
        logo(ws, "B20:C21")
        impressao(ws, "A1:G22", escala=88, ajustar_largura=0, quebras=(22,))
        oculta = wb.create_sheet("Listas")
    else:
        ws = wb.create_sheet("Edificações")
        larguras(ws, {"A": 3.0, "B": 30.0, "C": 14.0, "D": 14.0, "E": 14.0})
        mesclar(ws, "B2:E2", "CUSTO DE REPRODUÇÃO DAS EDIFICAÇÕES (EXEMPLO)", negrito=True, fundo=CABECALHO)
        for n, nome in enumerate(["Edificação", "Área (m²)", "Custo unitário (R$/m²)", "Custo (R$)"]):
            mesclar(ws, f"{get_column_letter(2 + n)}4", nome, negrito=True, fundo=ROTULO)
        for linha in range(5, 12):
            ws[f"E{linha}"] = f'=IFERROR(C{linha}*D{linha},"")'
            for col in "BCDE":
                moldura(ws, f"{col}{linha}")
        impressao(ws, "A1:F14", escala=95, paisagem=True)
        oculta = wb.create_sheet("Tab_Calculos")
    oculta["A1"] = "Logo"  # aba oculta: não recebe a logo (não dá pra ativar nem imprime)
    for n, item in enumerate(["BOM", "REGULAR", "RUIM"], 2):
        oculta[f"A{n}"] = item
    oculta.sheet_state = "hidden"


def gerar_banco_c():
    pasta = os.path.join(REPO, "modelos", BANCO_C.pasta_modelos)
    for tipo in ("Comparativo", "Evolutivo"):
        wb = Workbook()
        _aba_resumo_banco_c(wb, tipo)
        _aba_locais_banco_c(wb)
        _aba_fotos_banco_c(wb)
        _abas_extras_banco_c(wb, tipo)
        wb.active = 0
        salvar(wb, os.path.join(pasta, f"modelo_{tipo.lower()}.xlsm"))


# ---------------------------------------------------------------- Banco D

TIPOS_BANCO_D = {  # arquivo -> (tipo de operação em FORMULARIO!F36, variação)
    "modelo_aquisição_simples.xlsx": ("AQUISIÇÃO SIMPLES", None),
    "modelo_aquisição_à_vista_FGTS.xlsx": ("AQUISIÇÃO A VISTA COM FGTS", None),
    "modelo_consórcio_urbano.xlsx": ("CONSÓRCIO (IMÓVEL URBANO)", None),
    "modelo_consórcio_reforma_sem_ampliação.xlsx": ("CONSÓRCIO - REFORMA S/AMPLIAÇÃO (IMÓVEL URBANO)", "reforma"),
    "modelo_garantia_imóvel_urbano.xlsx": ("GARANTIA (IMÓVEL URBANO)", "liquidez"),
}
F_DATA = "dd/mm/yyyy"
F_MOEDA = '"R$" #,##0.00'
F_AREA = "0.00"

# (linha, rótulo em E, faixa do campo ou None, formato)
FORMULARIO_CAMPOS = [
    (5, "Junção:", "F5", "0"), (6, "Nome da junção:", "F6:L6", None), (7, "Protocolo:", "F7:H7", "@"),
    (8, "Nº do laudo:", "F8:H8", None), (9, "Grupo (só consórcio):", "F9:H9", None),
    (10, "Cota (só consórcio):", "F10:H10", None), (11, "Proponente:", "F11:L11", None),
    (12, "Endereço (tipo / nome):", None, None), (13, "Nº:", "F13", None), (14, "Complemento:", "F14:L14", None),
    (15, "Bairro:", "F15:L15", None), (16, "Cidade:", "F16:L16", None), (17, "CEP:", "F17:G17", "00000-000"),
    (18, "UF:", "F18", None), (19, "Data de recebimento:", "F19:G19", F_DATA), (20, "Data da visita:", "F20:G20", F_DATA),
    (36, "Tipo de operação:", "F36:L36", None),
    (37, "Qtde. de vagas de garagem:", "F37", None), (38, "Possui vagas?", "F38", None),
    (42, "Terreno (m²):", "F42", F_AREA), (43, "Construída (m²):", "F43", F_AREA), (44, "Comum (m²):", "F44", F_AREA),
    (45, "Garagem (m²):", "F45", F_AREA), (46, "Depósito (m²):", "F46", F_AREA), (47, "Total (m²):", "F47", F_AREA),
    (48, "Fração ideal:", "F48", "0.000000"),
    (57, "Comercialização:", "F57:H57", None), (58, "Localização:", "F58:H58", None),
    (59, "Regime de ocupação:", "F59:H59", None), (60, "Comentário:", "F60:L60", None),
    (61, "Data do laudo:", "F61:G61", F_DATA), (62, "Engenheiro(a) responsável:", "F62:L62", None),
    (63, "CREA:", "F63:G63", "0000000000000"), (64, "Local:", "F64:H64", None),
    (65, "Imóvel residencial com uso comercial?", "F65:H65", None),
    (143, "O imóvel tem condições de habitabilidade?", "F143:H143", None),
    (144, "Fatores valorizantes:", "F144:L144", None), (145, "Fatores desvalorizantes:", "F145:L145", None),
    (147, "Foi verificado risco ambiental?", "F147:H147", None), (148, "Imóvel em construção?", "F148:H148", None),
    (155, "Uso do prédio:", "F155:H155", None), (156, "Qtde. de pavimentos:", "F156", None),
    (157, "Qtde. de subsolos:", "F157", None), (158, "Unidades por andar:", "F158", None),
    (159, "Qtde. de elevadores:", "F159", None), (160, "Conservação do prédio:", "F160:H160", None),
    (161, "Piscina:", "F161", None), (162, "Playground:", "F162", None), (163, "Salão de festas:", "F163", None),
    (164, "Quadra:", "F164", None), (165, "Churrasqueira:", "F165", None), (166, "Gerador:", "F166", None),
    (167, "Conservação do condomínio:", "F167:H167", None),
]
ITENS_RESTRITIVOS = ["Área contaminada", "Área de preservação ambiental", "Imóvel encravado",
                     "Imóvel tombado", "Sem infraestrutura básica", "Mais de uma matrícula",
                     "Área embargada", "Risco geológico"]
SERVICOS = ["Coleta de lixo", "Transporte coletivo", "Comércio", "Rede bancária", "Escola", "Saúde",
            "Segurança", "Lazer"]
INFRA = ["Água", "Esgoto sanitário", "Energia elétrica", "Telefone", "Pavimentação", "Arborização",
         "Esgoto pluvial", "Gás canalizado", "Iluminação pública", "Fossa", "Poço"]
LISTAS = {  # validação de dados (lista suspensa) de alguns campos, como no formulário real
    '"BOM,REGULAR,RUIM"': ["F57", "F58"],
    '"OCUPADO,DESOCUPADO"': ["F59"],
    '"SIM,NÃO "': ["F38"],
    '"SIM ,NÃO ,NÃO SE APLICA"': ["F143"],
    '"ALTO,NORMAL,BAIXO,MINIMO"': ["E110", "J75"],
    '"RESIDENCIAL,COMERCIAL,MISTO,INDUSTRIAL"': ["G112", "J73", "F155"],
}


def _formulario(wb, tipo, variacao):
    ws = wb.active
    ws.title = "FORMULARIO"
    larguras(ws, {"A": 3.0, "B": 3.0, "C": 12.0, "D": 4.0, "E": 48.0, "F": 11.0, "G": 13.0, "H": 13.0, "I": 11.0,
                  "J": 12.0, "K": 10.0, "L": 10.0, "M": 10.0, "N": 10.0, "O": 10.0, "P": 10.0})
    alturas(ws, {1: 40.0, 60: 90.0})
    mesclar(ws, "C1:P1", "FORMULÁRIO PADRÃO - BANCO D (MODELO DE EXEMPLO)", negrito=True, fundo=CABECALHO,
            tamanho=14)
    rotulo(ws, "E4", "Empresa avaliadora:")
    campo(ws, "F4:L4", valor="EMPRESA DE AVALIAÇÃO (EXEMPLO)")
    for linha, texto, faixa, formato in FORMULARIO_CAMPOS:
        rotulo(ws, f"E{linha}", texto)
        if faixa:
            campo(ws, faixa, formato)
    # campos que já chegam preenchidos com dados de outro caso no modelo do banco
    # (perfis: pre_preenchimento.limpar apaga antes de escrever)
    ws["F5"] = 9999
    ws["F6"] = "AGÊNCIA DE OUTRO CASO"
    ws["F36"] = tipo
    campo(ws, "F12")
    campo(ws, "G12:L12")
    ws["F12"].number_format = "@"
    # itens restritivos (22-33): marca "X" na coluna G (esquerda) ou M (direita)
    rotulo(ws, "E22", "Foi assinalado algum item restritivo?", negrito=True)
    for n, item in enumerate(ITENS_RESTRITIVOS):
        linha, (col_txt, col_x) = 23 + n % 4, (("E", "G") if n < 4 else ("I", "M"))
        rotulo(ws, f"{col_txt}{linha}", item)
        campo(ws, f"{col_x}{linha}")
    ws["F21"] = "=COUNTA(G22:G33)+COUNTA(M22:M33)"
    # matrícula (33-34)
    for col, texto in (("F", "CNS Cartório"), ("H", "Livro"), ("J", "Nº Matrícula"), ("L", "Dígito")):
        rotulo(ws, f"{col}33", texto, negrito=True)
        campo(ws, f"{col}34", "@")
    rotulo(ws, "E34", "Matrícula do imóvel:")
    rotulo(ws, "M37", "Descobertas:")
    campo(ws, "N37")
    rotulo(ws, "O37", "Cobertas:")
    campo(ws, "P37")
    rotulo(ws, "E40", "ÁREAS - MATRÍCULA", negrito=True)
    rotulo(ws, "I40", "ESTIMADA", negrito=True)
    for linha, texto in ((42, "Terreno:"), (43, "Construída:")):
        rotulo(ws, f"I{linha}", texto)
        campo(ws, f"J{linha}", F_AREA)
    rotulo(ws, "I44", "IPTU", negrito=True)
    for linha, texto in ((45, "Terreno:"), (46, "Construída:")):
        rotulo(ws, f"I{linha}", texto)
        campo(ws, f"J{linha}", F_AREA)
    rotulo(ws, "I47", "Áreas sem averbação (m²):")
    campo(ws, "I48", F_AREA)
    ws["F47"] = "=F44+F43"
    # região (69-78)
    rotulo(ws, "E67", "CARACTERÍSTICAS DA REGIÃO", negrito=True)
    rotulo(ws, "E69", "Distância até (m):")
    for col, texto in (("F", "500"), ("G", "1000"), ("H", ">1000")):
        rotulo(ws, f"{col}69", texto, negrito=True)
    for n, servico in enumerate(SERVICOS):
        rotulo(ws, f"E{71 + n}", servico + ":")
        for col in "FGH":
            campo(ws, f"{col}{71 + n}")
    for linha, texto in ((70, "Bairro"), (72, "Predominância"), (74, "Padrão das construções"), (76, "Tráfego")):
        rotulo(ws, f"J{linha}", texto, negrito=True)
        campo(ws, f"J{linha + 1}:L{linha + 1}")
    # terreno (80-101)
    rotulo(ws, "E80", "TERRENO", negrito=True)
    rotulo(ws, "E82", "Infraestrutura urbana", negrito=True)
    for n, item in enumerate(INFRA):
        rotulo(ws, f"E{83 + n}", item)
        campo(ws, f"F{83 + n}")
    for linha, texto, formato in ((82, "Área (m²)", F_AREA), (85, "Zoneamento", None), (88, "Topografia", None),
                                  (91, "Formato", None)):
        rotulo(ws, f"I{linha}", texto, negrito=True)
        campo(ws, f"I{linha + 1}:J{linha + 1}", formato)
    rotulo(ws, "E95", "Dimensões (m)", negrito=True)
    for n, texto in enumerate(["Frente:", "Fundos:", "Lado direito:", "Lado esquerdo:"]):
        rotulo(ws, f"E{96 + n}", texto)
        campo(ws, f"F{96 + n}", "0.00")
    rotulo(ws, "E100", "Coordenada geográfica:", negrito=True)
    campo(ws, "E101")
    # imóvel (103-114): rótulo numa linha, campo na de baixo
    rotulo(ws, "E103", "IMÓVEL", negrito=True)
    for linha, esq, dir_ in ((105, "Tipo do imóvel", "Estrutura"), (107, "Tipo de implantação", "Fechamento das paredes"),
                             (109, "Padrão de acabamento", "Esquadrias"), (111, "Estado de conservação", "Uso do imóvel"),
                             (113, "Idade aparente (anos)", "Nº de pavimentos")):
        rotulo(ws, f"E{linha}", esq, negrito=True)
        rotulo(ws, f"G{linha}", dir_, negrito=True)
        campo(ws, f"E{linha + 1}")
        campo(ws, f"G{linha + 1}:I{linha + 1}")
    # cômodos (116-130): 3 linhas fixas + 10 livres; a PAG.2 mostra as 13
    rotulo(ws, "E115", "DESCRIÇÃO DO IMÓVEL", negrito=True)
    for col, texto in (("E", "Compartimento"), ("H", "Qtde."), ("I", "Paredes"), ("J", "Pisos"), ("K", "Tetos"),
                       ("L", "Pintura")):
        mesclar(ws, f"{col}117", texto, negrito=True, fundo=ROTULO)
    for linha in range(118, 131):
        for col in "EHIJKL":
            moldura(ws, f"{col}{linha}")
    ws["E118"], ws["E119"], ws["E120"] = "Quartos", "Suítes (quarto + banheiro)", "Banheiro"
    # valores (133-139) e vistoriador (135-138)
    for linha, texto in ((133, "Valor final de avaliação:"), (135, "Valor de liquidez:"), (137, "Valor de locação:"),
                         (139, "Valor de operação:")):
        rotulo(ws, f"G{linha}", texto, negrito=True)
        campo(ws, f"J{linha}:K{linha}", F_MOEDA)
    ws["J139"] = "=J133"
    for linha, texto, formato in ((135, "Vistoriador:", None), (136, "CREA/CAU:", None), (137, "Data:", F_DATA),
                                  (138, "Hora:", "h:mm")):
        rotulo(ws, f"C{linha}", texto)
        campo(ws, f"E{linha}", formato)
    rotulo(ws, "E153", "CARACTERÍSTICAS DO PRÉDIO (apartamento)", negrito=True)
    area_impressao = "C1:P171"
    if variacao == "reforma":  # serviços de reforma (aparecem na PAG.1, linhas 61-64)
        for linha, texto in ((170, "Descrição dos serviços de reforma:"), (171, "Custo da reforma:"),
                             (172, "Prazo para conclusão:")):
            rotulo(ws, f"E{linha}", texto)
            campo(ws, f"F{linha}:L{linha}")
        area_impressao = "C1:P182"
    for formula, celulas in LISTAS.items():
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        for c in celulas:
            dv.add(c)
        ws.add_data_validation(dv)
    impressao(ws, area_impressao, escala=40, quebras=(65, 132, 141))


def _pag1(wb, variacao):
    """PAG.1: os 3 quadros de perfis.BANCO_D.locais com o texto dentro, o
    resumo (fórmulas pra FORMULARIO) e o quadro da assinatura."""
    ws = wb.create_sheet(BANCO_D.locais.nome)
    larguras(ws, {"A": 12.57, **{get_column_letter(c): 9.14 for c in range(2, 12)}})
    # todas as linhas com altura fixa (nenhuma cresce sozinha): a página cabe numa folha a 66%
    alturas(ws, {linha: 15.0 for linha in range(1, 69)})
    alturas(ws, {3: 6.95, 5: 60.0, 22: 6.95, 28: 21.0, 29: 27.75, 30: 6.95, 31: 21.0, 32: 34.5,
                 33: 6.95, 34: 9.95, 35: 19.5, 36: 6.95, 50: 6.95, 67: 29.25, 68: 29.25})
    logo(ws)
    mesclar(ws, "C1:K2", "Laudo de avaliação - Banco D (modelo de exemplo)", negrito=True, fundo=CABECALHO,
            tamanho=12)
    for celula, texto in (("A4", "Junção:"), ("F4", "Protocolo:"), ("J4", "Nº laudo:")):
        rotulo(ws, celula, texto)
    ws["B4"], ws["C4"], ws["H4"], ws["K4"] = ("=+FORMULARIO!F5", "=+FORMULARIO!F6", "=+FORMULARIO!F7",
                                              "=+FORMULARIO!F8")
    mesclar(ws, "A5:K5", borda=False)
    ws["A5"].font = Font(bold=True, color="C00000", size=12)
    for q in BANCO_D.locais.quadros:  # texto dentro do quadro: sai quando a foto entra (limpar_texto)
        mesclar(ws, q.bloco, q.titulo, fundo=ROTULO)
    # identificação
    for celula, texto, formula in (("A23", "Grupo:", ("B23", "=+FORMULARIO!F9")), ("C23", "Cota:", ("D23", "=+FORMULARIO!F10")),
                                   ("A24", "Proponente:", ("B24", "=+FORMULARIO!F11")),
                                   ("A25", "Endereço:", ("B25", '=+FORMULARIO!F12&" "&FORMULARIO!G12')),
                                   ("H25", "Nº:", ("I25", "=+FORMULARIO!F13")),
                                   ("A26", "Compl.:", ("B26", "=+FORMULARIO!F14")),
                                   ("D26", "Bairro:", ("E26", "=+FORMULARIO!F15")),
                                   ("A27", "Cidade:", ("B27", "=+FORMULARIO!F16")),
                                   ("E27", "CEP:", ("F27", "=+FORMULARIO!F17")), ("H27", "UF:", ("I27", "=+FORMULARIO!F18")),
                                   ("J24", "Recebimento:", ("J25", "=+FORMULARIO!F19")),
                                   ("J26", "Data da visita:", ("J27", "=+FORMULARIO!F20"))):
        rotulo(ws, celula, texto)
        ws[formula[0]] = formula[1]
    ws["F27"].number_format = "00000-000"
    ws["J25"].number_format = ws["J27"].number_format = F_DATA
    mesclar(ws, "A28:K28", "Tipo de operação", negrito=True, fundo=CABECALHO)
    mesclar(ws, "A29:K29", "=+FORMULARIO!F36")
    if variacao == "liquidez":  # Garantia: valor de liquidez e aviso automático de item restritivo
        mesclar(ws, "A31:K31", "Valor de Liquidez", negrito=True, fundo=CABECALHO)
        mesclar(ws, "A32:D32", '=+IF(A31="Valor de Mercado",FORMULARIO!J133,FORMULARIO!J135)')
        ws["R5"] = "VIDE CARTA COM PARECER DO BANCO D \n(conforme normativo interno)"
        ws["A5"] = '=IF(COUNTIF(FORMULARIO!G22:G33,"X")+COUNTIF(FORMULARIO!M22:M33,"X")>0,R5,"")'
    else:
        mesclar(ws, "A31:K31", "Valor de Mercado", negrito=True, fundo=CABECALHO)
        mesclar(ws, "A32:D32", "=FORMULARIO!J139")
    ws["A32"].number_format = F_MOEDA
    rotulo(ws, "E32", "(valor por extenso)")
    for col, texto, ref in (("C", "CNS Cartório", "F34"), ("D", "Livro", "H34"), ("E", "Nº Matrícula", "J34"),
                            ("F", "Dígito", "L34")):
        rotulo(ws, f"{col}34", texto, negrito=True)
        ws[f"{col}35"] = f"=FORMULARIO!{ref}"
    rotulo(ws, "A35", "Matrícula:")
    mesclar(ws, "A37:F37", "Áreas (m²)", negrito=True, fundo=CABECALHO)
    mesclar(ws, "G37:K37", "Geral", negrito=True, fundo=CABECALHO)
    for n, (texto, ref) in enumerate((("Terreno:", "F42"), ("Construída:", "F43"), ("Comum:", "F44"),
                                      ("Total:", "F47"), ("Fração ideal:", "F48"), ("Sem averbação:", "I48"))):
        rotulo(ws, f"A{38 + n}", texto)
        ws[f"B{38 + n}"] = f"=+FORMULARIO!{ref}"
        ws[f"B{38 + n}"].number_format = F_AREA
    for n, (texto, formula) in enumerate((("Tipo do imóvel:", "=FORMULARIO!E106"),
                                          ("Idade aparente:", "=FORMULARIO!E114"),
                                          ("Quartos:", "=+SUM(FORMULARIO!H118:H119)"),
                                          ("Comercialização:", "=FORMULARIO!F57"),
                                          ("Ocupação:", "=+FORMULARIO!F59"),
                                          ("Uso comercial?", "=FORMULARIO!F65"))):
        rotulo(ws, f"G{38 + n}", texto)
        ws[f"J{38 + n}"] = formula
    mesclar(ws, "A47:K47", "Vagas de garagem", negrito=True, fundo=CABECALHO)
    rotulo(ws, "A48", "Qtde. de vagas:")
    ws["C48"] = "=+FORMULARIO!F37"
    rotulo(ws, "A49", "Possui vagas?")
    ws["C49"] = "=+FORMULARIO!F38"
    mesclar(ws, "A51:K51", "Comentários", negrito=True, fundo=CABECALHO)
    if variacao == "reforma":  # comentário até a 60; 61-64 mostram a reforma
        mesclar(ws, "A52:K60", "=+FORMULARIO!F60", alinhar=TEXTO_LONGO)
        for faixa, formula in (("A61:L61", "=FORMULARIO!E170"), ("A62:L62", "=FORMULARIO!F170"),
                               ("A63:B63", "=FORMULARIO!E171"), ("C63:D63", "=FORMULARIO!F171"),
                               ("A64:B64", "=FORMULARIO!E172"), ("C64:D64", "=FORMULARIO!F172")):
            mesclar(ws, faixa, formula, alinhar=ESQUERDA, borda=False)
    else:
        mesclar(ws, "A52:K64", "=+FORMULARIO!F60", alinhar=TEXTO_LONGO)
    for celula, texto, ref, formato in (("A65", "Vistoriador:", "B65", None), ("F65", "Data:", "G65", F_DATA),
                                        ("A66", "CREA/CAU:", "B66", None), ("F66", "Hora:", "G66", "h:mm"),
                                        ("A67", "Local:", "B67", None), ("H67", "Engº.:", "I67", None),
                                        ("A68", "Data do laudo:", "B68", F_DATA), ("H68", "CREA:", "I68", None)):
        rotulo(ws, celula, texto)
        ws[ref] = {"B65": "=+FORMULARIO!E135", "G65": "=+FORMULARIO!E137", "B66": "=+FORMULARIO!E136",
                   "G66": "=+FORMULARIO!E138", "B67": "=+FORMULARIO!F64", "I67": "=+FORMULARIO!F62",
                   "B68": "=+FORMULARIO!F61", "I68": "=+FORMULARIO!F63"}[ref]
        if formato:
            ws[ref].number_format = formato
    mesclar(ws, BANCO_D.assinatura.area)  # quadro da assinatura da engenheira responsável
    # escala fixa (não "ajustar à página"): as páginas de FOTOS COMPLEMENTARES
    # criadas depois da linha 68 mantêm o tamanho (ver perfis.BANCO_D)
    impressao(ws, "A1:K68", escala=59 if variacao == "reforma" else 66)
    ws.print_options.horizontalCentered = True


def _pag2(wb):
    """PAG.2: região, terreno e as 13 linhas de cômodos (FORMULARIO 118-130)."""
    ws = wb.create_sheet("PAG.2")
    larguras(ws, {"A": 11.14})
    logo(ws)
    mesclar(ws, "C1:L2", "Laudo de avaliação - Banco D (modelo de exemplo)", negrito=True, fundo=CABECALHO)
    mesclar(ws, "A6:L6", "Características da região", negrito=True, fundo=CABECALHO)
    for n, servico in enumerate(SERVICOS):
        rotulo(ws, f"A{8 + n}", servico + ":")
        for col, ref in zip("CEG", "FGH"):
            ws[f"{col}{8 + n}"] = f"=+FORMULARIO!{ref}{71 + n}"
    mesclar(ws, "A18:L18", "Terreno e imóvel", negrito=True, fundo=CABECALHO)
    for n, item in enumerate(INFRA):
        ws[f"A{20 + n}"] = f"=+FORMULARIO!F{83 + n}"
        rotulo(ws, f"B{20 + n}", item)
    mesclar(ws, "A35:L35", "Descrição do imóvel", negrito=True, fundo=CABECALHO)
    for col, texto in (("A", "Compartimento"), ("E", "Qtde."), ("F", "Paredes"), ("H", "Pisos"), ("I", "Tetos"),
                       ("K", "Pintura")):
        rotulo(ws, f"{col}38", texto, negrito=True)
    for n in range(13):  # PAG.2!A47 = FORMULARIO!E126 (no modelo original era 0 - ver montar_modelos_banco_d.py)
        for col, ref in zip("AEFHIK", "EHIJKL"):
            ws[f"{col}{39 + n}"] = f"=+FORMULARIO!{ref}{118 + n}"
    mesclar(ws, "A53:L53", "Características do prédio", negrito=True, fundo=CABECALHO)
    for n in range(13):
        ws[f"H{55 + n}"] = f"=FORMULARIO!F{155 + n}"
    impressao(ws, "A1:L68", escala=68)
    ws.print_options.horizontalCentered = True


def _croqui(wb):
    ws = wb.create_sheet("CROQUI-PLANTA")
    logo(ws)
    mesclar(ws, "C1:L2", "Croqui / planta do imóvel", negrito=True, fundo=CABECALHO)
    mesclar(ws, "A4:L4", "(desenho feito pela equipe)", borda=False)
    impressao(ws, "A1:L17", escala=84)


def _fotos_banco_d(wb):
    """FOTOS: só o cabeçalho e um texto de instrução (A7:L14) - a grade é
    montada pelo programa (perfis.Grade)."""
    lay = BANCO_D.fotos
    ws = wb.create_sheet(lay.nome)
    larguras(ws, {"A": 9.14, "J": 12.14, "K": 9.14, "L": 7.29})
    alturas(ws, {3: 6.95, 5: 6.95, 6: 21.0, 9: 21.0, 10: 21.0, 11: 21.0, 12: 21.0, 13: 21.0, 14: 21.0})
    logo(ws)
    mesclar(ws, "C1:L2", "Laudo de avaliação - Banco D (modelo de exemplo)", negrito=True, fundo=CABECALHO)
    for celula, texto, formula in (("A4", "Junção:", ("B4", "=+PAG.1!B4")), ("F4", "Protocolo:", ("H4", "=+PAG.1!H4")),
                                   ("J4", "Nº laudo:", ("K4", "=+PAG.1!K4"))):
        rotulo(ws, celula, texto)
        ws[formula[0]] = formula[1]
    mesclar(ws, f"A6:{lay.col_fim_impressao}6", "RELATÓRIO FOTOGRÁFICO", negrito=True, fundo=CABECALHO)
    for celula, texto in (("B9", "Inserir as fotos no formato que achar conveniente."), ("B11", "Lembrete:"),
                          ("B12", "* fotos de todos os ambientes;"), ("B13", "* fotos do logradouro.")):
        rotulo(ws, celula, texto, negrito=True)
    impressao(ws, "A1:L23", escala=82)


def gerar_banco_d():
    pasta = os.path.join(REPO, "modelos", BANCO_D.pasta_modelos)
    for arquivo, (tipo, variacao) in TIPOS_BANCO_D.items():
        wb = Workbook()
        _formulario(wb, tipo, variacao)
        _pag1(wb, variacao)
        _pag2(wb)
        _croqui(wb)
        _fotos_banco_d(wb)
        wb.active = 0
        salvar(wb, os.path.join(pasta, arquivo))


if __name__ == "__main__":
    gerar_banco_c()
    gerar_banco_d()
